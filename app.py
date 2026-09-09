import streamlit as st
import pandas as pd
from datetime import date
from io import BytesIO
from PIL import Image as PilImage
import os

# Importaciones de ReportLab (PDF)
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as RLImage, Table, TableStyle
from reportlab.lib.units import inch
from reportlab.lib import colors

# Importaciones de python-docx (Word)
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Importación de traductor gratuito (Sin API Key)
from deep_translator import GoogleTranslator

# --- 1. CONFIGURACIÓN Y PERSISTENCIA ---
st.set_page_config(page_title="Generador de Reportes", page_icon="🔧", layout="wide")

# Inicialización de estados para los 6 documentos
claves_estado = ['pdf_es', 'word_es', 'txt_es', 'pdf_en', 'word_en', 'txt_en']
for key in claves_estado:
    if key not in st.session_state:
        st.session_state[key] = None

# --- 2. CARGA DE DATOS ---
@st.cache_data
def cargar_datos_servicios():
    if os.path.exists("servicios.csv"):
        for encoding in ['utf-8', 'latin-1', 'iso-8859-1', 'cp1252']:
            try:
                df = pd.read_csv("servicios.csv", dtype=str, encoding=encoding, sep=None, engine='python')
                df.columns = df.columns.str.strip()
                nombres_clave = {'Serie/Artículo': 'Serie', 'Fec. Fac. Min': 'Fec_Fac_Min', 'Fac. Min': 'Fac_Min'}
                df = df.rename(columns=nombres_clave)
                return df
            except:
                continue
    return pd.DataFrame(columns=['Orden', 'Cliente', 'Serie', 'Producto', 'Fec_Fac_Min', 'Fac_Min'])

df_db = cargar_datos_servicios()

# --- CONSTANTES ---
LISTA_TECNICOS = [
    "Tec. Xavier Ramón", "Tec. Juan Diego Quezada", "Tec. Javier Quiguango",
    "Tec. Wilson Quiguango", "Tec. Carlos Jama", "Tec. Manuel Vera",
    "Tec. Juan Farez", "Tec. Santiago Farez", "Tec.Ronald Paladinez", "Tec. Saul Vite",
    "Tec. Miguel Meza", "Tec German Tenemaza", "CISTRONIC", "Taller Externo"
]
LISTA_REALIZADORES = ["Ing. Henry Beltran", "Ing. Pablo Lopez", "Ing. Christian Calle", "Ing. Guillermo Ortiz", "Ing. Adrian Aguilar", "Ing. John Juela"]
OPCIONES_REPORTE = ["FUERA DE GARANTIA", "INFORME TECNICO", "RECLAMO AL PROVEEDOR"]

TEXTOS_CONCLUSIONES = {
    "FUERA DE GARANTIA": "En marco de las políticas de garantía que mantienen un orden en el proceso se concluye:\nEl daño identificado no es atribuible a defectos de fabricación o materiales, sino al uso indebido del equipo, lo cual invalida la cobertura de garantía.",
    "INFORME TECNICO": "En marco de las políticas de garantía que mantienen un orden en el proceso se concluye:\nCon base en estos hallazgos se informa que el equipo funciona correctamente en base a las indicaciones de operacion del fabricante",
    "RECLAMO AL PROVEEDOR": "En marco de las políticas de garantía que mantienen un orden en el proceso se concluye:\nEl daño presentado corresponde a un defecto de manufactura debido al diagnostico realizado. Solicitamos su colaboración con el proceso de Reclamo al Proveedor."
}

# --- 3. FUNCIONES DE TRADUCCIÓN ---
def traducir_texto(texto, idioma):
    if idioma == "Español" or not texto.strip():
        return texto
    try:
        traductor = GoogleTranslator(source='es', target='en')
        return traductor.translate(texto)
    except Exception as e:
        return f"[Error de traducción] {texto}"

# Etiquetas estáticas bilingües
LBL = {
    "Español": {
        "titulo": "INFORME TÉCNICO DE SERVICIO", "orden": "Orden", "fac": "Factura", "cliente": "Cliente",
        "fec_fac": "Fec. Factura", "prod": "Producto", "serie": "Serie", "realizador": "Realizado por",
        "fec_rep": "Fecha Reporte", "revisador": "Revisado por", "evidencia": "EVIDENCIA DE IMÁGENES",
        "figura": "Figura"
    },
    "Inglés": {
        "titulo": "TECHNICAL SERVICE REPORT", "orden": "Order", "fac": "Invoice", "cliente": "Customer",
        "fec_fac": "Invoice Date", "prod": "Product", "serie": "Serial/Code", "realizador": "Prepared by",
        "fec_rep": "Report Date", "revisador": "Reviewed by", "evidencia": "IMAGE EVIDENCE",
        "figura": "Figure"
    }
}

# --- 4. FUNCIONES DE GENERACIÓN ---
def agregar_marca_agua(canvas, doc):
    watermark_file = "watermark.png"
    if os.path.exists(watermark_file):
        try:
            canvas.saveState()
            canvas.setFillAlpha(0.12)
            canvas.drawImage(watermark_file, 0, 0, width=canvas._pagesize[0], height=canvas._pagesize[1], mask='auto', preserveAspectRatio=True, anchor='c')
            canvas.restoreState()
        except:
            pass

def generar_pdf(datos, secciones_activas, lista_imgs, idioma):
    l = LBL[idioma]
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=0.4*inch, bottomMargin=0.4*inch, leftMargin=0.5*inch, rightMargin=0.5*inch)
    color_azul = colors.HexColor("#0056b3")
    
    est_titulo = ParagraphStyle('T', fontSize=16, alignment=1, fontName='Helvetica-Bold', textColor=color_azul)
    est_sec = ParagraphStyle('S', fontSize=10, fontName='Helvetica-Bold', textColor=colors.white, backColor=color_azul, borderPadding=2, spaceBefore=8)
    est_txt = ParagraphStyle('TXT', fontSize=9, fontName='Helvetica', leading=11)
    est_firma = ParagraphStyle('F', fontSize=10, fontName='Helvetica-Bold', alignment=1)
    est_fig = ParagraphStyle('FIG', fontSize=10, fontName='Helvetica-Bold', alignment=1, spaceBefore=4)
    
    story = []

    # Cabecera (Logos)
    logo_izq_path, logo_der_path = "logo.png", "logo_derecho.png"
    col_izq, col_der = [], []
    if os.path.exists(logo_izq_path): col_izq.append(RLImage(logo_izq_path, width=1.4*inch, height=0.55*inch))
    if os.path.exists(logo_der_path):
        img_der = RLImage(logo_der_path, width=1.4*inch, height=0.55*inch)
        img_der.hAlign = 'RIGHT'
        col_der.append(img_der)

    if col_izq or col_der:
        header_table = Table([[col_izq, col_der]], colWidths=[3.7*inch, 3.7*inch])
        header_table.setStyle(TableStyle([('VALIGN', (0,0), (-1,-1), 'TOP'), ('ALIGN', (1,0), (1,0), 'RIGHT')]))
        story.append(header_table)
    
    story.append(Spacer(1, 10))
    story.append(Paragraph(l['titulo'], est_titulo))
    story.append(Spacer(1, 15))
    
    fac_txt = "STOCK" if str(datos['factura']).strip() in ["0", "nan", ""] else datos['factura']
    info = [
        [Paragraph(f"<b>{l['orden']}:</b> {datos['orden']}", est_txt), Paragraph(f"<b>{l['fac']}:</b> {fac_txt}", est_txt)],
        [Paragraph(f"<b>{l['cliente']}:</b> {datos['cliente']}", est_txt), Paragraph(f"<b>{l['fec_fac']}:</b> {datos['fecha_factura']}", est_txt)],
        [Paragraph(f"<b>{l['prod']}:</b> {datos['producto']}", est_txt), Paragraph(f"<b>{l['serie']}:</b> {datos['serie']}", est_txt)],
        [Paragraph(f"<b>{l['realizador']}:</b> {datos['realizador']}", est_txt), Paragraph(f"<b>{l['fec_rep']}:</b> {datos['fecha_reporte']}", est_txt)]
    ]
    t = Table(info, colWidths=[3.7*inch, 3.7*inch])
    t.setStyle(TableStyle([('GRID', (0,0), (-1,-1), 0.5, colors.lightgrey), ('VALIGN', (0,0), (-1,-1), 'MIDDLE')]))
    story.append(t)
    story.append(Spacer(1, 10))

    # Secciones Dinámicas
    for tit, cont in secciones_activas:
        story.append(Paragraph(tit, est_sec))
        story.append(Paragraph(cont.replace('\n', '<br/>'), est_txt))
        story.append(Spacer(1, 5))

    # Imágenes
    if lista_imgs:
        story.append(Paragraph(l['evidencia'], est_sec))
        story.append(Spacer(1, 10))
        datos_cuadricula, fila_actual = [], []
        
        for idx, i in enumerate(lista_imgs):
            try:
                img_stream = BytesIO(i['imagen_raw'])
                img_obj = RLImage(img_stream, width=3.4*inch, height=2.2*inch)
                celda = [img_obj, Spacer(1, 4), Paragraph(f"{l['figura']} {idx+1}. {i['descripcion']}", est_fig)]
                fila_actual.append(celda)
                if len(fila_actual) == 2:
                    datos_cuadricula.append(fila_actual)
                    fila_actual = []
            except Exception:
                fila_actual.append([Paragraph(f"Error imagen {idx+1}", est_txt)])
                
        if fila_actual:
            fila_actual.append("")
            datos_cuadricula.append(fila_actual)
            
        if datos_cuadricula:
            t_grid = Table(datos_cuadricula, colWidths=[3.7*inch, 3.7*inch])
            t_grid.setStyle(TableStyle([
                ('GRID', (0,0), (-1,-1), 1, colors.black),
                ('VALIGN', (0,0), (-1,-1), 'MIDDLE'), ('ALIGN', (0,0), (-1,-1), 'CENTER'),
                ('PADDING', (0,0), (-1,-1), 6)
            ]))
            story.append(t_grid)

    story.append(Spacer(1, 60))
    t_firmas = Table([[Paragraph(f"{l['realizador']}:", est_firma), Paragraph(f"{l['revisador']}:", est_firma)], 
                      [Paragraph(datos['realizador'], est_firma), Paragraph(datos['tecnico'], est_firma)]], colWidths=[3.7*inch, 3.7*inch])
    story.append(t_firmas)
    
    doc.build(story, onFirstPage=agregar_marca_agua, onLaterPages=agregar_marca_agua)
    buffer.seek(0)
    return buffer.read()

def generar_word(datos, secciones_activas, lista_imgs, idioma):
    l = LBL[idioma]
    doc = Document()
    
    # --- INCORPORACIÓN DE LOGOS EN WORD ---
    logo_izq_path, logo_der_path = "logo.png", "logo_derecho.png"
    
    if os.path.exists(logo_izq_path) or os.path.exists(logo_der_path):
        # Crear tabla invisible de 1 fila y 2 columnas
        tabla_logos = doc.add_table(rows=1, cols=2)
        celda_izq = tabla_logos.cell(0, 0)
        celda_der = tabla_logos.cell(0, 1)
        
        if os.path.exists(logo_izq_path):
            p_izq = celda_izq.paragraphs[0]
            run_izq = p_izq.add_run()
            run_izq.add_picture(logo_izq_path, width=Inches(1.4))
            p_izq.alignment = WD_ALIGN_PARAGRAPH.LEFT
            
        if os.path.exists(logo_der_path):
            p_der = celda_der.paragraphs[0]
            run_der = p_der.add_run()
            run_der.add_picture(logo_der_path, width=Inches(1.4))
            p_der.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            
        doc.add_paragraph() # Espaciador debajo de los logos
    # --- FIN DE LOGOS ---

    h1 = doc.add_heading(l['titulo'], level=1)
    h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    fac_txt = "STOCK" if str(datos['factura']).strip() in ["0", "nan", ""] else datos['factura']
    table = doc.add_table(rows=4, cols=2)
    table.style = 'Table Grid'
    
    celdas = table.rows
    celdas[0].cells[0].text = f"{l['orden']}: {datos['orden']}"
    celdas[0].cells[1].text = f"{l['fac']}: {fac_txt}"
    celdas[1].cells[0].text = f"{l['cliente']}: {datos['cliente']}"
    celdas[1].cells[1].text = f"{l['fec_fac']}: {datos['fecha_factura']}"
    celdas[2].cells[0].text = f"{l['prod']}: {datos['producto']}"
    celdas[2].cells[1].text = f"{l['serie']}: {datos['serie']}"
    celdas[3].cells[0].text = f"{l['realizador']}: {datos['realizador']}"
    celdas[3].cells[1].text = f"{l['fec_rep']}: {datos['fecha_reporte']}"
    
    doc.add_paragraph()

    for tit, cont in secciones_activas:
        doc.add_heading(tit, level=2)
        doc.add_paragraph(cont)

    if lista_imgs:
        doc.add_heading(l['evidencia'], level=2)
        for idx, i in enumerate(lista_imgs):
            try:
                img_stream = BytesIO(i['imagen_raw'])
                doc.add_picture(img_stream, width=Inches(4.5))
                p = doc.add_paragraph(f"{l['figura']} {idx+1}. {i['descripcion']}")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            except Exception as e:
                doc.add_paragraph(f"Error cargando imagen {idx+1}: {e}")

    doc.add_paragraph("\n\n\n")
    p_firmas = doc.add_paragraph()
    p_firmas.add_run(f"{l['realizador']}:\t\t\t\t\t{l['revisador']}:\n").bold = True
    p_firmas.add_run(f"{datos['realizador']}\t\t\t\t\t{datos['tecnico']}")
    
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()

def generar_txt_contenido(datos, secciones_activas, idioma):
    l = LBL[idioma]
    fac_txt = "STOCK" if str(datos['factura']).strip() in ["0", "nan", ""] else datos['factura']
    
    txt = f"{l['titulo']}\n" + "="*30 + "\n\n"
    txt += f"{l['cliente']}: {datos['cliente']}\n"
    txt += f"{l['fac']}: {fac_txt}\n"
    txt += f"{l['fec_fac']}: {datos['fecha_factura']}\n"
    txt += f"{l['orden']}: {datos['orden']}\n"
    txt += f"{l['serie']}: {datos['serie']}\n"
    txt += f"{l['prod']}: {datos['producto']}\n"
    txt += f"{l['revisador']}: {datos['tecnico']}\n\n"
    txt += f"TIPO DE REPORTE: {datos['tipo_reporte']}\n\n"
    
    for tit, cont in secciones_activas:
        txt += f"--- {tit} ---\n{cont}\n\n"

    txt += f"Atentamente,\n{datos['realizador']}\nCoordinador Postventa"
    return txt

# --- 5. INTERFAZ ---
st.title("🚀 Gestión de Reportes Técnicos (Bilingüe)")

orden_id = st.text_input("Ingrese número de Orden")
c_v, s_v, p_v, f_v, ff_v = "", "", "", "", date.today()

if orden_id:
    res = df_db[df_db['Orden'] == orden_id]
    if not res.empty:
        row = res.iloc[0]
        c_v, s_v, p_v, f_v = str(row.get('Cliente','')), str(row.get('Serie','')), str(row.get('Producto','')), str(row.get('Fac_Min',''))
        try: ff_v = pd.to_datetime(str(row.get('Fec_Fac_Min',''))).date()
        except: ff_v = date.today()

st.markdown("### Datos del Reporte")
col1, col2 = st.columns(2)
with col1:
    tipo_rep = st.selectbox("Tipo de Reporte", options=OPCIONES_REPORTE)
    f_realizador = st.selectbox("Realizado por", options=LISTA_REALIZADORES)
    f_cliente = st.text_input("Cliente", value=c_v)
    f_prod = st.text_input("Producto", value=p_v)
with col2:
    f_tecnico = st.selectbox("Revisado por (Técnico)", options=LISTA_TECNICOS)
    f_fac = st.text_input("Factura", value=f_v)
    f_fec_fac = st.date_input("Fecha Factura", value=ff_v)
    f_fec_rep = st.date_input("Fecha del Reporte", value=date.today())
    f_serie = st.text_input("Serie/Artículo", value=s_v)

st.markdown("---")
st.markdown("### 📝 Contenido del Reporte (Selecciona qué incluir)")

def_rf = f"Ingresa a servicio técnico {f_prod}. Se observa el uso continuo del artículo."
def_re = "Se procede a revisar el sistema de alimentación de energía y sus líneas de conexión.\nSe procede a revisar el sistema electrónico del equipo.\nSe procede a revisar el sistema mecanico de equipo"
def_obs = "Luego de la revisión del artículo se observa lo siguiente: "
texto_concl_default = TEXTOS_CONCLUSIONES.get(tipo_rep, "")

c_chk1, c_txt1 = st.columns([1, 10])
with c_chk1: inc_rf = st.checkbox("Incluir", value=True, key="c1")
with c_txt1: f_rev_fisica = st.text_area("1. Revisión Física", value=def_rf)

c_chk2, c_txt2 = st.columns([1, 10])
with c_chk2: inc_ing = st.checkbox("Incluir", value=True, key="c2")
with c_txt2: f_ingreso_tec = st.text_area("2. Ingresa a servicio técnico")

c_chk3, c_txt3 = st.columns([1, 10])
with c_chk3: inc_re = st.checkbox("Incluir", value=True, key="c3")
with c_txt3: f_rev_electro = st.text_area("3. Revisión electro-electrónica-mecanica", value=def_re)

c_chk4, c_txt4 = st.columns([1, 10])
with c_chk4: inc_obs = st.checkbox("Incluir", value=True, key="c4")
with c_txt4: f_obs = st.text_area("4. Observaciones", value=def_obs)

c_chk5, c_txt5 = st.columns([1, 10])
with c_chk5: inc_con = st.checkbox("Incluir", value=True, key="c5")
with c_txt5: f_concl = st.text_area("5. Conclusiones", value=texto_concl_default, height=150)

st.markdown("---")
st.markdown("### 📸 Evidencia Fotográfica")
uploaded_files = st.file_uploader("Subir imágenes", type=['jpg','png','jpeg'], accept_multiple_files=True)

lista_imgs_temp = [] 
if uploaded_files:
    for idx, file in enumerate(uploaded_files):
        c_img, c_desc = st.columns([1, 3])
        with c_img:
            st.image(file, use_container_width=True)
        with c_desc:
            desc = st.text_input(f"Descripción Imagen #{idx+1}", value="Evidencia técnica.", key=f"desc_{idx}")
            lista_imgs_temp.append({"file": file, "desc": desc})

st.markdown("---")

if st.button("💾 GENERAR ARCHIVOS EN ESPAÑOL E INGLÉS", use_container_width=True):
    with st.spinner("Generando los 6 documentos simultáneamente..."):
        
        imgs_procesadas_es = []
        imgs_procesadas_en = []
        for item in lista_imgs_temp:
            try:
                p_img = PilImage.open(item['file'])
                if p_img.mode != 'RGB': p_img = p_img.convert('RGB')
                img_byte = BytesIO()
                p_img.save(img_byte, format='JPEG', quality=95)
                img_raw = img_byte.getvalue() 
                
                desc_es = item['desc']
                desc_en = traducir_texto(desc_es, "Inglés")
                
                imgs_procesadas_es.append({"imagen_raw": img_raw, "descripcion": desc_es})
                imgs_procesadas_en.append({"imagen_raw": img_raw, "descripcion": desc_en})
            except Exception as e:
                st.error(f"Error procesando imagen: {e}")

        # --- DATOS EN ESPAÑOL ---
        secciones_es = []
        titulos_es = ["1. Revisión Física", "2. Ingresa a servicio técnico", "3. Revisión electro-electrónica-mecanica", "4. Observaciones", "5. Conclusiones"]
        
        if inc_rf: secciones_es.append((titulos_es[0], f_rev_fisica))
        if inc_ing: secciones_es.append((titulos_es[1], f_ingreso_tec))
        if inc_re: secciones_es.append((titulos_es[2], f_rev_electro))
        if inc_obs: secciones_es.append((titulos_es[3], f_obs))
        if inc_con: secciones_es.append((titulos_es[4], f_concl))

        datos_es = {
            "orden": orden_id, "cliente": f_cliente, "factura": f_fac, "fecha_factura": f_fec_fac,
            "producto": f_prod, "serie": f_serie, "tecnico": f_tecnico, "realizador": f_realizador,
            "fecha_reporte": f_fec_rep, "tipo_reporte": tipo_rep
        }

        # --- DATOS EN INGLÉS ---
        secciones_en = []
        titulos_en = ["1. Physical Inspection", "2. Entry to Technical Service", "3. Electro-Mechanical Review", "4. Observations", "5. Conclusions"]
        
        if inc_rf: secciones_en.append((titulos_en[0], traducir_texto(f_rev_fisica, "Inglés")))
        if inc_ing: secciones_en.append((titulos_en[1], traducir_texto(f_ingreso_tec, "Inglés")))
        if inc_re: secciones_en.append((titulos_en[2], traducir_texto(f_rev_electro, "Inglés")))
        if inc_obs: secciones_en.append((titulos_en[3], traducir_texto(f_obs, "Inglés")))
        if inc_con: secciones_en.append((titulos_en[4], traducir_texto(f_concl, "Inglés")))

        datos_en = {
            "orden": orden_id, "cliente": f_cliente, "factura": f_fac, "fecha_factura": f_fec_fac,
            "producto": traducir_texto(f_prod, "Inglés"), "serie": f_serie, "tecnico": f_tecnico, "realizador": f_realizador,
            "fecha_reporte": f_fec_rep, "tipo_reporte": traducir_texto(tipo_rep, "Inglés")
        }

        st.session_state.pdf_es = generar_pdf(datos_es, secciones_es, imgs_procesadas_es, "Español")
        st.session_state.word_es = generar_word(datos_es, secciones_es, imgs_procesadas_es, "Español")
        st.session_state.txt_es = generar_txt_contenido(datos_es, secciones_es, "Español")
        
        st.session_state.pdf_en = generar_pdf(datos_en, secciones_en, imgs_procesadas_en, "Inglés")
        st.session_state.word_en = generar_word(datos_en, secciones_en, imgs_procesadas_en, "Inglés")
        st.session_state.txt_en = generar_txt_contenido(datos_en, secciones_en, "Inglés")
        
        st.success("✅ ¡Los 6 archivos fueron generados exitosamente!")

if st.session_state.pdf_es:
    st.markdown("### 🇪🇸 Descargas en Español")
    c1, c2, c3 = st.columns(3)
    with c1:
        st.download_button("📄 Descargar PDF (ES)", data=st.session_state.pdf_es, file_name=f"{f_prod}_{f_cliente}_{orden_id}_ES.pdf", mime="application/pdf", use_container_width=True)
    with c2:
        st.download_button("📝 Descargar Word (ES)", data=st.session_state.word_es, file_name=f"{f_prod}_{f_cliente}_{orden_id}_ES.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", use_container_width=True)
    with c3:
        st.download_button("🔤 Descargar TXT (ES)", data=st.session_state.txt_es, file_name=f"Status_{orden_id}_ES.txt", mime="text/plain", use_container_width=True)
        
    st.markdown("### 🇺🇸 Descargas en Inglés")
    c4, c5, c6 = st.columns(3)
    with c4:
        st.download_button("📄 Descargar PDF (EN)", data=st.session_state.pdf_en, file_name=f"{f_prod}_{f_cliente}_{orden_id}_EN.pdf", mime="application/pdf", use_container_width=True)
    with c5:
        st.download_button("📝 Descargar Word (EN)", data=st.session_state.word_en, file_name=f"{f_prod}_{f_cliente}_{orden_id}_EN.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", use_container_width=True)
    with c6:
        st.download_button("🔤 Descargar TXT (EN)", data=st.session_state.txt_en, file_name=f"Status_{orden_id}_EN.txt", mime="text/plain", use_container_width=True)
